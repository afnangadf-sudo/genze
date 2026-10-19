#!/usr/bin/env python3
"""Generate binary test fixtures for Step 3 file-upload tests."""
import os

HERE = os.path.dirname(os.path.abspath(__file__))

# --- minimal valid PDF with one text line (pdf.js-parseable) ---
text = "Genze PDF fixture - the quick brown fox jumps over the lazy dog."
stream = f"BT /F1 18 Tf 72 720 Td ({text}) Tj ET".encode("latin-1")
objs = []
objs.append(b"<< /Type /Catalog /Pages 2 0 R >>")
objs.append(b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>")
objs.append(b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>")
objs.append(b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream")
objs.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
pdf = [b"%PDF-1.4\n"]
offsets = []
for i, body in enumerate(objs, start=1):
    offsets.append(sum(len(p) for p in pdf))
    pdf.append(f"{i} 0 obj\n".encode() + body + b"\nendobj\n")
xref_pos = sum(len(p) for p in pdf)
pdf.append(f"xref\n0 {len(objs)+1}\n".encode())
pdf.append(b"0000000000 65535 f \n")
for off in offsets:
    pdf.append(f"{off:010d} 00000 n \n".encode())
pdf.append(f"trailer\n<< /Size {len(objs)+1} /Root 1 0 R >>\nstartxref\n{xref_pos}\n%%EOF".encode())
with open(os.path.join(HERE, "hello.pdf"), "wb") as f:
    f.write(b"".join(pdf))

# --- DOCX via python-docx ---
from docx import Document
doc = Document()
doc.add_paragraph("Genze DOCX fixture - pack my box with five dozen liquor jugs.")
doc.add_paragraph("Second paragraph for good measure.")
doc.save(os.path.join(HERE, "hello.docx"))

# --- XLSX via openpyxl ---
from openpyxl import Workbook
wb = Workbook()
ws = wb.active
ws.title = "Planets"
ws.append(["Name", "Moons", "Note"])
ws.append(["Earth", 1, "Genze XLSX fixture"])
ws.append(["Mars", 2, "the red planet"])
ws2 = wb.create_sheet("Empty")
ws2.append(["only", "header"])
wb.save(os.path.join(HERE, "hello.xlsx"))

# --- CSV ---
with open(os.path.join(HERE, "hello.csv"), "w") as f:
    f.write("city,country,note\nLahore,Pakistan,Genze CSV fixture\nParis,France,croissants\n")

# --- PNG via PIL ---
from PIL import Image
img = Image.new("RGB", (320, 180), color=(190, 120, 60))
img.save(os.path.join(HERE, "hello.png"))

# --- MP3: dummy bytes (content never parsed without a key; mime is what matters) ---
with open(os.path.join(HERE, "hello.mp3"), "wb") as f:
    f.write(b"ID3" + b"\x00" * 4096)

print("fixtures written to", HERE)
